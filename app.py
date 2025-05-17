import os
import unicodedata
import streamlit as st
from dotenv import load_dotenv
from PyPDF2 import PdfReader
import requests
from fpdf import FPDF
from io import BytesIO
from docx import Document

# ── Init & Config ──────────────────────────────────────────────────────────────
load_dotenv()
GROQ_API_KEY = os.getenv("GROQ_API_KEY")
if not GROQ_API_KEY:
    st.error("❌ GROQ_API_KEY missing in .env. Add it and restart.")
    st.stop()

st.set_page_config(page_title="CV & Cover Letter Generator", layout="centered")
st.title("🎯 CV & Cover Letter Generator (Groq Cloud)")

# Initialize session state
if "cv_out" not in st.session_state:
    st.session_state.cv_out = None
if "cover_out" not in st.session_state:
    st.session_state.cover_out = None

# ── UI Inputs ───────────────────────────────────────────────────────────────────
resume_file = st.file_uploader("📄 Upload your resume (PDF)", type=["pdf"])
jd = st.text_area(
    "🧐 Paste the Job Description here:",
    height=200,
    placeholder="e.g. “Seeking an experienced Python developer with AWS…”"
)

# ── Helpers ─────────────────────────────────────────────────────────────────────
def parse_resume(pdf_file) -> str:
    reader = PdfReader(pdf_file)
    return "\n\n".join(page.extract_text() or "" for page in reader.pages)

def generate_documents(resume_text: str, jd_text: str) -> tuple[str, str]:
    prompt = (
        "You are a career coach. Given the candidate’s resume below and the job description, produce:\n\n"
        "1) A complete, one-page professional CV formatted with Markdown headings:\n"
        "   ## Name & Contact Information\n"
        "   ## Professional Summary\n"
        "   ## Key Skills\n"
        "   ## Work Experience\n"
        "   ## Education & Certifications\n\n"
        "2) A crisp, professional cover letter addressed “Dear Hiring Manager,” tailored to the role.\n\n"
        "=== Resume ===\n"
        f"{resume_text}\n\n"
        "=== Job Description ===\n"
        f"{jd_text}\n\n"
        "Return the CV first, then on its own line “---”, then the cover letter."
    )
    r = requests.post(
        "https://api.groq.com/openai/v1/chat/completions",
        headers={
            "Content-Type": "application/json",
            "Authorization": f"Bearer {GROQ_API_KEY}"
        },
        json={
            "model": "llama-3.3-70b-versatile",
            "messages": [{"role": "user", "content": prompt}],
            "temperature": 0.7,
            "max_tokens": 1500
        }
    )
    r.raise_for_status()
    content = r.json()["choices"][0]["message"]["content"]
    cv_text, cover_text = content.split("---", 1)
    return cv_text.strip(), cover_text.strip()

def render_pdf(text: str) -> bytes:
    # Replace bullets with dashes, normalize to ASCII
    text = text.replace("•", "-")
    safe = (
        unicodedata.normalize("NFKD", text)
                   .encode("ascii", "ignore")
                   .decode("ascii")
    )
    pdf = FPDF()
    pdf.set_auto_page_break(True, margin=15)
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 16)

    for line in safe.split("\n"):
        if line.startswith("## "):
            pdf.ln(4)
            pdf.set_font("Helvetica", "B", 14)
            pdf.cell(0, 8, line[3:], ln=True)
            pdf.ln(1)
            pdf.set_font("Helvetica", "", 12)
        elif line.startswith("### "):
            pdf.set_font("Helvetica", "B", 12)
            pdf.cell(0, 7, line[4:], ln=True)
            pdf.ln(1)
            pdf.set_font("Helvetica", "", 12)
        elif line.strip().startswith("- "):
            pdf.cell(5)
            pdf.multi_cell(0, 6, "- " + line.strip()[2:])
        elif not line.strip():
            pdf.ln(3)
        else:
            pdf.multi_cell(0, 6, line)

    return pdf.output(dest="S").encode("latin-1")

def render_docx(text: str) -> bytes:
    doc = Document()
    for line in text.split("\n"):
        if line.strip() == "":
            doc.add_paragraph()
        else:
            doc.add_paragraph(line)
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()

# ── Generate Action ────────────────────────────────────────────────────────────
if st.button("🚀 Generate"):
    if not resume_file or not jd.strip():
        st.warning("⚠️ Please upload a resume PDF and paste a job description.")
    else:
        with st.spinner("🔍 Parsing resume…"):
            resume_text = parse_resume(resume_file)
        with st.spinner("✍️ Generating…"):
            try:
                cv, cover = generate_documents(resume_text, jd)
            except requests.HTTPError as e:
                st.error(f"❌ API request failed: {e}")
            else:
                # Store in session state so it persists
                st.session_state.cv_out = cv
                st.session_state.cover_out = cover

# ── Display Generated & Download ───────────────────────────────────────────────
if st.session_state.cv_out:
    st.subheader("📝 Your One-Page CV")
    st.write(st.session_state.cv_out)
    pdf_cv = render_pdf(st.session_state.cv_out)
    docx_cv = render_docx(st.session_state.cv_out)
    st.download_button(
        "📥 Download CV (PDF)",
        pdf_cv,
        "custom_cv.pdf",
        "application/pdf"
    )
    st.download_button(
        "📥 Download CV (DOCX)",
        docx_cv,
        "custom_cv.docx",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )

if st.session_state.cover_out:
    st.subheader("📄 Your Cover Letter")
    st.write(st.session_state.cover_out)
    pdf_cl = render_pdf(st.session_state.cover_out)
    docx_cl = render_docx(st.session_state.cover_out)
    st.download_button(
        "📥 Download Cover Letter (PDF)",
        pdf_cl,
        "cover_letter.pdf",
        "application/pdf"
    )
    st.download_button(
        "📥 Download Cover Letter (DOCX)",
        docx_cl,
        "cover_letter.docx",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
